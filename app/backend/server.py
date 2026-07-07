from fastapi import FastAPI, APIRouter, HTTPException, Depends, Header, Query
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import io
import csv
import logging
import uuid
import bcrypt
import secrets
from pathlib import Path
from pydantic import BaseModel, Field, ConfigDict
from typing import List, Optional, Literal
from datetime import datetime, timezone, date
from openpyxl import Workbook

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

# ---------------- Preset data ----------------
PRESET_EXPENSE_CATEGORIES = [
    "Food and Drinks", "Travel", "Stationary", "Salary Paid", "Reference Commission",
    "Fuel", "Bills", "Maintenance", "Online Shopping", "Offline Shopping",
    "Household", "Insurance", "Training and Software", "Fees", "Loans",
    "Gifts", "Vehicle Servicing", "Investments", "Medical", "Internet & Telecommunication"
]
PRESET_INCOME_CATEGORIES = [
    "LIC of India", "NJ Funds", "Star Health", "Niva Bupa", "HDFC Ergo",
    "HDFC Life", "Tata AIA Life Insurance", "Prudent", "Advisory", "Turtlemint"
]

# ---------------- Models ----------------
class Entry(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    date: str
    amount: float
    type: Literal["income", "expense"]
    scope: Literal["personal", "business"]
    category: str
    note: Optional[str] = ""
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class EntryCreate(BaseModel):
    date: str
    amount: float
    type: Literal["income", "expense"]
    scope: Literal["personal", "business"]
    category: str
    note: Optional[str] = ""

class EntryUpdate(BaseModel):
    date: Optional[str] = None
    amount: Optional[float] = None
    type: Optional[Literal["income", "expense"]] = None
    scope: Optional[Literal["personal", "business"]] = None
    category: Optional[str] = None
    note: Optional[str] = None

class Category(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    type: Literal["income", "expense"]
    is_preset: bool = False

class CategoryCreate(BaseModel):
    name: str
    type: Literal["income", "expense"]

class PinSetup(BaseModel):
    pin: str

class PinVerify(BaseModel):
    pin: str

# ---------------- Auth (PIN) ----------------
def _hash_pin(pin: str) -> str:
    return bcrypt.hashpw(pin.encode(), bcrypt.gensalt()).decode()

def _verify_pin(pin: str, hashed: str) -> bool:
    return bcrypt.checkpw(pin.encode(), hashed.encode())

async def require_auth(x_session_token: Optional[str] = Header(None)):
    if not x_session_token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    sess = await db.sessions.find_one({"token": x_session_token}, {"_id": 0})
    if not sess:
        raise HTTPException(status_code=401, detail="Invalid or expired session")
    return True

@api_router.get("/auth/status")
async def auth_status():
    pin_doc = await db.auth.find_one({"key": "pin"}, {"_id": 0})
    return {"pin_set": pin_doc is not None}

@api_router.post("/auth/setup")
async def auth_setup(body: PinSetup):
    if not (body.pin.isdigit() and 4 <= len(body.pin) <= 8):
        raise HTTPException(status_code=400, detail="PIN must be 4-8 digits")
    existing = await db.auth.find_one({"key": "pin"})
    if existing:
        raise HTTPException(status_code=400, detail="PIN already set. Use reset.")
    await db.auth.insert_one({"key": "pin", "hash": _hash_pin(body.pin)})
    token = secrets.token_urlsafe(32)
    await db.sessions.insert_one({"token": token, "created_at": datetime.now(timezone.utc).isoformat()})
    return {"token": token}

@api_router.post("/auth/verify")
async def auth_verify(body: PinVerify):
    pin_doc = await db.auth.find_one({"key": "pin"}, {"_id": 0})
    if not pin_doc:
        raise HTTPException(status_code=400, detail="PIN not set")
    if not _verify_pin(body.pin, pin_doc["hash"]):
        raise HTTPException(status_code=401, detail="Incorrect PIN")
    token = secrets.token_urlsafe(32)
    await db.sessions.insert_one({"token": token, "created_at": datetime.now(timezone.utc).isoformat()})
    return {"token": token}

@api_router.post("/auth/logout")
async def auth_logout(x_session_token: Optional[str] = Header(None)):
    if x_session_token:
        await db.sessions.delete_one({"token": x_session_token})
    return {"ok": True}

# ---------------- Categories ----------------
@api_router.get("/categories", response_model=List[Category])
async def list_categories(_: bool = Depends(require_auth)):
    docs = await db.categories.find({}, {"_id": 0}).to_list(1000)
    return [Category(**d) for d in docs]

@api_router.post("/categories", response_model=Category)
async def add_category(body: CategoryCreate, _: bool = Depends(require_auth)):
    existing = await db.categories.find_one({"name": body.name, "type": body.type})
    if existing:
        raise HTTPException(status_code=400, detail="Category already exists")
    cat = Category(name=body.name, type=body.type, is_preset=False)
    await db.categories.insert_one(cat.model_dump())
    return cat

@api_router.delete("/categories/{cat_id}")
async def delete_category(cat_id: str, _: bool = Depends(require_auth)):
    doc = await db.categories.find_one({"id": cat_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    if doc.get("is_preset"):
        raise HTTPException(status_code=400, detail="Cannot delete preset category")
    await db.categories.delete_one({"id": cat_id})
    return {"ok": True}

# ---------------- Entries ----------------
@api_router.get("/entries", response_model=List[Entry])
async def list_entries(
    _: bool = Depends(require_auth),
    year: Optional[int] = None,
    month: Optional[int] = None,
    fy_start: Optional[int] = None,
    scope: Optional[str] = None,
    type: Optional[str] = None,
):
    q = {}
    if year and month:
        prefix = f"{year:04d}-{month:02d}"
        q["date"] = {"$regex": f"^{prefix}"}
    elif year:
        q["date"] = {"$regex": f"^{year:04d}"}
    elif fy_start:
        q["$or"] = [
            {"date": {"$gte": f"{fy_start:04d}-04-01", "$lte": f"{fy_start:04d}-12-31"}},
            {"date": {"$gte": f"{fy_start+1:04d}-01-01", "$lte": f"{fy_start+1:04d}-03-31"}},
        ]
    if scope:
        q["scope"] = scope
    if type:
        q["type"] = type
    docs = await db.entries.find(q, {"_id": 0}).sort("date", -1).to_list(10000)
    return [Entry(**d) for d in docs]

@api_router.post("/entries", response_model=Entry)
async def create_entry(body: EntryCreate, _: bool = Depends(require_auth)):
    entry = Entry(**body.model_dump())
    await db.entries.insert_one(entry.model_dump())
    return entry

@api_router.put("/entries/{entry_id}", response_model=Entry)
async def update_entry(entry_id: str, body: EntryUpdate, _: bool = Depends(require_auth)):
    existing = await db.entries.find_one({"id": entry_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    updates = {k: v for k, v in body.model_dump().items() if v is not None}
    await db.entries.update_one({"id": entry_id}, {"$set": updates})
    doc = await db.entries.find_one({"id": entry_id}, {"_id": 0})
    return Entry(**doc)

@api_router.delete("/entries/{entry_id}")
async def delete_entry(entry_id: str, _: bool = Depends(require_auth)):
    res = await db.entries.delete_one({"id": entry_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Not found")
    return {"ok": True}

# ---------------- Summaries ----------------
def _fy_bounds(fy_start: int):
    return f"{fy_start:04d}-04-01", f"{fy_start+1:04d}-03-31"

async def _fetch_range(start: str, end: str):
    docs = await db.entries.find(
        {"date": {"$gte": start, "$lte": end}},
        {"_id": 0}
    ).to_list(100000)
    return docs

@api_router.get("/summary/monthly")
async def monthly_summary(_: bool = Depends(require_auth), year: int = Query(...), month: int = Query(...)):
    prefix = f"{year:04d}-{month:02d}"
    docs = await db.entries.find({"date": {"$regex": f"^{prefix}"}}, {"_id": 0}).to_list(10000)
    result = {
        "personal": {"income": 0.0, "expense": 0.0, "by_category": {}},
        "business": {"income": 0.0, "expense": 0.0, "by_category": {}},
    }
    for d in docs:
        s = result[d["scope"]]
        s[d["type"]] += d["amount"]
        key = d["category"]
        s["by_category"].setdefault(key, {"income": 0.0, "expense": 0.0})
        s["by_category"][key][d["type"]] += d["amount"]
    for scope in ("personal", "business"):
        r = result[scope]
        r["net"] = r["income"] - r["expense"]
    totals = {
        "income": result["personal"]["income"] + result["business"]["income"],
        "expense": result["personal"]["expense"] + result["business"]["expense"],
    }
    totals["net"] = totals["income"] - totals["expense"]
    result["totals"] = totals
    return result

@api_router.get("/summary/yearly")
async def yearly_summary(_: bool = Depends(require_auth), fy_start: int = Query(...)):
    """Indian FY: fy_start-04-01 to (fy_start+1)-03-31, 12 months."""
    start, end = _fy_bounds(fy_start)
    docs = await _fetch_range(start, end)
    months = []
    for i in range(12):
        m = ((3 + i) % 12) + 1
        y = fy_start if i < 9 else fy_start + 1
        months.append({"year": y, "month": m, "label": f"{y:04d}-{m:02d}"})

    per_month = {mm["label"]: {
        "personal_income": 0.0, "personal_expense": 0.0,
        "business_income": 0.0, "business_expense": 0.0,
    } for mm in months}

    for d in docs:
        label = d["date"][:7]
        if label not in per_month:
            continue
        key = f"{d['scope']}_{d['type']}"
        per_month[label][key] += d["amount"]

    rows = []
    for mm in months:
        row = {**mm, **per_month[mm["label"]]}
        row["personal_net"] = row["personal_income"] - row["personal_expense"]
        row["business_net"] = row["business_income"] - row["business_expense"]
        row["total_income"] = row["personal_income"] + row["business_income"]
        row["total_expense"] = row["personal_expense"] + row["business_expense"]
        row["total_net"] = row["total_income"] - row["total_expense"]
        rows.append(row)

    totals = {
        "personal_income": sum(r["personal_income"] for r in rows),
        "personal_expense": sum(r["personal_expense"] for r in rows),
        "business_income": sum(r["business_income"] for r in rows),
        "business_expense": sum(r["business_expense"] for r in rows),
    }
    totals["personal_net"] = totals["personal_income"] - totals["personal_expense"]
    totals["business_net"] = totals["business_income"] - totals["business_expense"]
    totals["total_income"] = totals["personal_income"] + totals["business_income"]
    totals["total_expense"] = totals["personal_expense"] + totals["business_expense"]
    totals["total_net"] = totals["total_income"] - totals["total_expense"]

    return {"fy_start": fy_start, "fy_label": f"FY {fy_start}-{str(fy_start+1)[-2:]}", "rows": rows, "totals": totals}

# ---------------- Export ----------------
async def _fetch_for_export(fy_start: Optional[int], year: Optional[int], month: Optional[int], scope: Optional[str]):
    q = {}
    if year and month:
        q["date"] = {"$regex": f"^{year:04d}-{month:02d}"}
    elif year:
        q["date"] = {"$regex": f"^{year:04d}"}
    elif fy_start:
        q["$or"] = [
            {"date": {"$gte": f"{fy_start:04d}-04-01", "$lte": f"{fy_start:04d}-12-31"}},
            {"date": {"$gte": f"{fy_start+1:04d}-01-01", "$lte": f"{fy_start+1:04d}-03-31"}},
        ]
    if scope:
        q["scope"] = scope
    return await db.entries.find(q, {"_id": 0}).sort("date", 1).to_list(100000)

@api_router.get("/export/csv")
async def export_csv(
    _: bool = Depends(require_auth),
    fy_start: Optional[int] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
    scope: Optional[str] = None,
):
    docs = await _fetch_for_export(fy_start, year, month, scope)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Date", "Scope", "Type", "Category", "Amount (INR)", "Note"])
    for d in docs:
        w.writerow([d["date"], d["scope"], d["type"], d["category"], f'{d["amount"]:.2f}', d.get("note", "")])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=finance_export.csv"},
    )

@api_router.get("/export/xlsx")
async def export_xlsx(
    _: bool = Depends(require_auth),
    fy_start: Optional[int] = None,
    year: Optional[int] = None,
    month: Optional[int] = None,
    scope: Optional[str] = None,
):
    docs = await _fetch_for_export(fy_start, year, month, scope)
    wb = Workbook()
    ws = wb.active
    ws.title = "Entries"
    headers = ["Date", "Scope", "Type", "Category", "Amount (INR)", "Note"]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="09090B")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")
    for d in docs:
        ws.append([d["date"], d["scope"], d["type"], d["category"], round(d["amount"], 2), d.get("note", "")])
    for col_letter, width in zip(["A", "B", "C", "D", "E", "F"], [14, 12, 12, 28, 16, 40]):
        ws.column_dimensions[col_letter].width = width

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=finance_export.xlsx"},
    )

# ---------------- Seed ----------------
async def seed_categories():
    count = await db.categories.count_documents({})
    if count > 0:
        return
    docs = []
    for name in PRESET_EXPENSE_CATEGORIES:
        docs.append(Category(name=name, type="expense", is_preset=True).model_dump())
    for name in PRESET_INCOME_CATEGORIES:
        docs.append(Category(name=name, type="income", is_preset=True).model_dump())
    await db.categories.insert_many(docs)

async def seed_sample_entries():
    count = await db.entries.count_documents({})
    if count > 0:
        return
    today = date.today()
    def d(offset):
        from datetime import timedelta
        return (today - timedelta(days=offset)).isoformat()
    samples = [
        {"date": d(1), "amount": 45000, "type": "income", "scope": "business", "category": "LIC of India", "note": "Policy commission"},
        {"date": d(2), "amount": 1250, "type": "expense", "scope": "personal", "category": "Food and Drinks", "note": "Groceries"},
        {"date": d(4), "amount": 22000, "type": "income", "scope": "business", "category": "NJ Funds", "note": "Trail commission"},
        {"date": d(6), "amount": 3200, "type": "expense", "scope": "personal", "category": "Fuel", "note": "Petrol"},
        {"date": d(9), "amount": 8500, "type": "expense", "scope": "business", "category": "Training and Software", "note": "CRM subscription"},
    ]
    entries = [Entry(**s).model_dump() for s in samples]
    await db.entries.insert_many(entries)

@app.on_event("startup")
async def startup_event():
    await seed_categories()
    await seed_sample_entries()

# ---------------- Root ----------------
@api_router.get("/")
async def root():
    return {"message": "Finance Tracker API"}

app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()